from __future__ import annotations
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document as WordDocument

def create_pdf(text: str, output_path: Path, title: str = "Clean Notes") -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=48,
        leftMargin=48,
        topMargin=48,
        bottomMargin=48,
    )
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for block in [b.strip() for b in text.split("\n") if b.strip()]:
        safe = (
            block.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        story.append(Paragraph(safe, styles["BodyText"]))
        story.append(Spacer(1, 6))
    doc.build(story)

def create_docx(text: str, output_path: Path, title: str = "Clean Notes") -> None:
    doc = WordDocument()
    doc.add_heading(title, level=1)
    for line in [x.strip() for x in text.split("\n") if x.strip()]:
        if line.startswith(("-", "•")):
            doc.add_paragraph(line.lstrip("-• ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(str(output_path))
